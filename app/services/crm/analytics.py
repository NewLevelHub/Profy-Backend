"""The owner's dashboard (ТЗ §3.12.5) and the analytics workbook for the head
of sales (ТЗ §3.12.6, П1-21).

Dashboard, over a period of Almaty days: each specialist's load (hours opened
and taken), the calls held, missed and cancelled, and the money of their
deals; the cancelled, moved and missed calls with the reasons; how long a
paid client waits for the review.

Workbook, for a month cohort (deals created in it, as on the «Аналитика»
page): the funnel, conversions by source and by manager. Ads by creative
come with the Meta Marketing API; until then the sheet says so. No personal
data in it, so the marketer gets the same file without the managers.
"""

import io
import statistics
import uuid
from collections import defaultdict
from datetime import date

from fastapi import status
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.i18n import DEFAULT_LOCALE
from app.i18n.catalog import tr
from app.models.crm import STAGE_ORDER, CrmDeal, DealResult
from app.models.scheduling import Meeting, MeetingStatus, MeetingType, SlotState, SpecialistSlot
from app.models.user import User, UserRole
from app.schemas import crm as schemas
from app.services import audit, tz
from app.services.crm import sources
from app.services.crm.errors import crm_error

MAX_DASHBOARD_DAYS = 366
SLOT_HOURS = 0.5


def _bounds(date_from: date, date_to: date):
    if date_to < date_from or (date_to - date_from).days + 1 > MAX_DASHBOARD_DAYS:
        raise crm_error(status.HTTP_422_UNPROCESSABLE_ENTITY, "period_invalid")
    return tz.day_bounds(date_from)[0], tz.day_bounds(date_to)[1]


def _name(user: User | None) -> str:
    return (user.full_name or user.email) if user else "—"


async def dashboard(db: AsyncSession, date_from: date, date_to: date) -> schemas.OwnerDashboard:
    start, end = _bounds(date_from, date_to)
    in_period = and_(Meeting.start_at >= start, Meeting.start_at < end)

    specialists = {
        u.id: u
        for u in (await db.execute(select(User).where(User.role == UserRole.psychologist))).scalars()
    }
    slots: dict[uuid.UUID, dict[SlotState, int]] = defaultdict(dict)
    for spec_id, state, count in await db.execute(
        select(SpecialistSlot.specialist_id, SpecialistSlot.state, func.count())
        .where(SpecialistSlot.start_at >= start, SpecialistSlot.start_at < end)
        .group_by(SpecialistSlot.specialist_id, SpecialistSlot.state)
    ):
        slots[spec_id][state] = count

    meetings = list((await db.execute(select(Meeting).where(in_period).order_by(Meeting.start_at))).scalars())
    # A move cancels the call and books a new one pointing back at it.
    moved_to = {
        m.rescheduled_from_id: m.start_at
        for m in (
            await db.execute(select(Meeting).where(Meeting.rescheduled_from_id.in_([m.id for m in meetings])))
        ).scalars()
    } if meetings else {}
    counts: dict[uuid.UUID, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    issues: list[schemas.MeetingIssue] = []
    for m in meetings:
        if m.status == MeetingStatus.done:
            counts[m.specialist_id]["done"] += 1
            continue
        if m.status == MeetingStatus.no_show:
            kind = "no_show"
        elif m.status == MeetingStatus.cancelled:
            kind = "rescheduled" if m.id in moved_to else "cancelled"
        else:
            continue
        counts[m.specialist_id][kind] += 1
        issues.append(schemas.MeetingIssue(
            meeting_id=m.id, deal_id=m.deal_id, kind=kind, type=m.type.value, start=m.start_at,
            specialist=_name(specialists.get(m.specialist_id)),
            student_name=(m.handoff or {}).get("student_name"),
            reason=m.cancel_reason, moved_to=moved_to.get(m.id),
        ))

    revenue = dict((await db.execute(
        select(CrmDeal.specialist_id, func.coalesce(func.sum(CrmDeal.paid_amount), 0))
        .where(CrmDeal.specialist_id.is_not(None), CrmDeal.paid_at >= start, CrmDeal.paid_at < end)
        .group_by(CrmDeal.specialist_id)
    )).all())

    rows = []
    for spec_id, user in specialists.items():
        opened = slots.get(spec_id, {})
        row = schemas.SpecialistLoad(
            id=spec_id, name=_name(user), is_active=user.is_active,
            free_hours=opened.get(SlotState.free, 0) * SLOT_HOURS,
            busy_hours=opened.get(SlotState.booked, 0) * SLOT_HOURS,
            calls_done=counts[spec_id]["done"], no_shows=counts[spec_id]["no_show"],
            cancelled=counts[spec_id]["cancelled"] + counts[spec_id]["rescheduled"],
            revenue=int(revenue.get(spec_id, 0)),
        )
        # Nothing in the period and no longer working: not a row.
        if user.is_active or row.free_hours or row.busy_hours or row.calls_done or row.revenue:
            rows.append(row)
    rows.sort(key=lambda r: (-r.busy_hours, r.name))

    return schemas.OwnerDashboard(
        from_=date_from, to=date_to, specialists=rows, issues=issues,
        payment_to_review=await _payment_to_review(db, start, end),
    )


async def _payment_to_review(db: AsyncSession, start, end) -> schemas.PaymentToReview:
    held = (await db.execute(
        select(CrmDeal.paid_at, func.min(Meeting.start_at))
        .join(Meeting, Meeting.deal_id == CrmDeal.id)
        .where(Meeting.type == MeetingType.main60, Meeting.status == MeetingStatus.done, CrmDeal.paid_at.is_not(None))
        .group_by(CrmDeal.id, CrmDeal.paid_at)
    )).all()
    days = [
        max((review - paid).total_seconds(), 0) / 86400
        for paid, review in held
        if start <= review < end
    ]
    has_review = select(Meeting.deal_id).where(
        Meeting.type == MeetingType.main60, Meeting.status == MeetingStatus.done, Meeting.deal_id.is_not(None)
    )
    waiting = await db.scalar(
        select(func.count()).select_from(CrmDeal).where(
            CrmDeal.paid_at.is_not(None),
            CrmDeal.psychologist_review.is_(True),
            CrmDeal.result == DealResult.open,
            CrmDeal.id.not_in(has_review),
        )
    )
    return schemas.PaymentToReview(
        count=len(days),
        median_days=round(statistics.median(days), 1) if days else None,
        average_days=round(statistics.fmean(days), 1) if days else None,
        max_days=round(max(days), 1) if days else None,
        waiting=waiting or 0,
    )


# ── workbook ────────────────────────────────────────────────────────────────


def _month_bounds(month: str):
    year, mon = (int(part) for part in month.split("-"))
    first = date(year, mon, 1)
    after = date(year + (mon == 12), mon % 12 + 1, 1)
    return tz.day_bounds(first)[0], tz.day_bounds(after)[0]


def _group_rows(deals: list[CrmDeal], key_of) -> list[list]:
    groups: dict[str, list[int]] = {}
    for deal in deals:
        row = groups.setdefault(key_of(deal), [0, 0, 0, 0, 0, 0])
        row[0] += 1
        row[1] += deal.is_target is True
        if deal.paid_at is not None:
            row[2] += 1
            row[3] += deal.paid_amount or 0
        row[4] += deal.result == DealResult.won
        row[5] += deal.result == DealResult.lost
    return sorted(([key, *values] for key, values in groups.items()), key=lambda r: -r[1])


def _share(part: int, whole: int) -> float | None:
    return round(part / whole, 4) if whole else None


def _sheet(book: Workbook, title: str, header: list[str], rows: list[list], percent_cols: tuple[int, ...] = ()):
    sheet = book.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for col in percent_cols:
        for (cell,) in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = "0%"
    for i, title_text in enumerate(header, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = max(12, len(title_text) + 2)
    sheet.column_dimensions["A"].width = 28
    sheet.freeze_panes = "A2"
    return sheet


def _conversion_rows(groups: list[list], total_label: str) -> list[list]:
    def line(key, leads, target, paid, revenue, won, lost):
        return [key, leads, target, _share(target, leads), paid, _share(paid, leads), revenue, won, lost]

    rows = [line(*g) for g in groups]
    if groups:
        rows.append(line(total_label, *(sum(g[i] for g in groups) for i in range(1, 7))))
    return rows


async def workbook(db: AsyncSession, user: User, month: str) -> bytes:
    labels = tr("crm_export", locale=DEFAULT_LOCALE)
    start, end = _month_bounds(month)
    deals = list((await db.execute(
        select(CrmDeal).where(CrmDeal.created_at >= start, CrmDeal.created_at < end)
    )).scalars())
    source_names = await sources.names(db)
    managers = {
        u.id: _name(u)
        for u in (await db.execute(select(User).where(User.id.in_({d.manager_id for d in deals if d.manager_id})))).scalars()
    }

    book = Workbook()
    book.remove(book.active)

    funnel = []
    previous = None
    for index, stage in enumerate(STAGE_ORDER):
        reached = sum(1 for d in deals if d.result == DealResult.won or STAGE_ORDER.index(d.stage) >= index)
        funnel.append([labels["stages"][stage.value], reached, _share(reached, len(deals)),
                       _share(reached, previous) if previous is not None else None])
        previous = reached
    sheet = _sheet(book, labels["sheet_funnel"], [labels["stage"], labels["reached"], labels["share"], labels["step_conversion"]], funnel, (3, 4))
    sheet.append([])
    sheet.append([labels["period"], month])
    sheet.append([labels["cohort_note"]])

    header = [labels["leads"], labels["target"], labels["target_share"], labels["paid"], labels["paid_conversion"],
              labels["revenue"], labels["won"], labels["lost"]]
    by_source = _group_rows(deals, lambda d: source_names.get(d.source_id) or labels["no_source"])
    _sheet(book, labels["sheet_sources"], [labels["source"], *header], _conversion_rows(by_source, labels["total"]), (4, 6))
    if user.role != UserRole.marketer:
        by_manager = _group_rows(deals, lambda d: managers.get(d.manager_id) or labels["no_manager"])
        _sheet(book, labels["sheet_managers"], [labels["manager"], *header], _conversion_rows(by_manager, labels["total"]), (4, 6))
    ads = _sheet(book, labels["sheet_ads"], [labels["source"], *header], _conversion_rows(by_source, labels["total"]), (4, 6))
    ads.insert_rows(1)
    ads["A1"] = labels["ads_pending"]
    ads.freeze_panes = "A3"

    audit.record(
        db, actor=user, action="crm.analytics_exported", entity_type="crm_analytics", entity_id=month,
        after={"deals": len(deals)},
    )
    await db.commit()
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
