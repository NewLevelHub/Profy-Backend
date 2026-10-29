"""Meta ads in the CRM (app/services/meta_ads/, app/services/crm/ads.py):
the sync upserts, the report counts the cost of a lead and a client from
the deals, the owner's rate turns USD into ₸. Meta is never called."""

import io
from datetime import timedelta
from unittest.mock import AsyncMock

import httpx
import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models.crm import CrmDeal
from app.models.meta_ads import MetaAd, MetaAdInsight, MetaFxRate, MetaSyncState
from app.models.user import UserRole
from app.services import tz
from app.services.meta_ads import client, sync
from tests.integration.test_crm_deals import URL, _create, _h, _user

TODAY = tz.today()
NAMES = {"a1": "Статика 1.4", "a2": "Видео 2"}


def _insight(day, ad_id: str, spend: str, conversations: int, **extra) -> dict:
    return {
        "date_start": day.isoformat(), "campaign_id": "c1", "campaign_name": "Переписки в WhatsApp",
        "adset_id": "s1", "adset_name": "Advantage+", "ad_id": ad_id, "ad_name": NAMES[ad_id],
        "spend": spend, "impressions": "1000", "reach": "800", "clicks": "30", "inline_link_clicks": "20",
        "actions": [
            {"action_type": "link_click", "value": "20"},
            {"action_type": "onsite_conversion.messaging_conversation_started_7d", "value": str(conversations)},
        ],
        **extra,
    }


@pytest_asyncio.fixture
async def meta(monkeypatch: pytest.MonkeyPatch, db_session: AsyncSession) -> dict[str, AsyncMock]:
    """Meta configured, its answers stubbed; ads data of this database
    dropped for the test (rolled back after)."""
    for model in (MetaAdInsight, MetaFxRate, MetaSyncState):
        await db_session.execute(delete(model))
    await db_session.execute(delete(MetaAd))
    # Kept past the sync's own rollback on an error (the outer test
    # transaction still undoes it).
    await db_session.commit()
    monkeypatch.setattr(settings, "META_ACCESS_TOKEN", "token")
    monkeypatch.setattr(settings, "META_AD_ACCOUNT_ID", "act_1")
    ads = AsyncMock(return_value=[{
        "id": "a1", "name": "Статика 1.4", "status": "ACTIVE", "adset_id": "s1", "campaign_id": "c1",
        "adset": {"name": "Advantage+"}, "campaign": {"name": "Переписки в WhatsApp"}, "creative": {"id": "cr1"},
    }])
    insights = AsyncMock(return_value=[
        _insight(TODAY, "a1", "10.00", 4),
        _insight(TODAY - timedelta(days=1), "a1", "5.50", 2),
        _insight(TODAY, "a2", "4.50", 1),
    ])
    monkeypatch.setattr(client, "ads", ads)
    monkeypatch.setattr(client, "insights", insights)
    return {"ads": ads, "insights": insights}


async def test_the_sync_upserts_and_pulls_less_after_the_first_run(db_session: AsyncSession, meta) -> None:
    assert await sync.sync(db_session) == 3
    since, until = meta["insights"].await_args.args
    assert until == TODAY and (until - since).days == sync.HISTORY_DAYS - 1

    ad = await db_session.get(MetaAd, "a1")
    assert (ad.name, ad.campaign_name, ad.creative_id) == ("Статика 1.4", "Переписки в WhatsApp", "cr1")
    assert (await db_session.get(MetaAd, "a2")).adset_name == "Advantage+"  # known from the insights only
    row = await db_session.get(MetaAdInsight, (TODAY, "a1"))
    assert (str(row.spend), row.conversations, row.link_clicks) == ("10.00", 4, 20)

    # Meta revised today: the same key, new numbers. The hourly run takes two days.
    meta["insights"].return_value = [_insight(TODAY, "a1", "12.40", 5)]
    await sync.sync(db_session)
    since, until = meta["insights"].await_args.args
    assert (until - since).days == 1 and meta["ads"].await_count == 1
    await db_session.refresh(row)
    assert (str(row.spend), row.conversations) == ("12.40", 5)
    assert await db_session.scalar(select(func.count()).select_from(MetaAdInsight)) == 3


async def test_a_revoked_token_is_recorded(db_session: AsyncSession, meta) -> None:
    meta["insights"].side_effect = client.MetaError(190, "Error validating access token")
    with pytest.raises(client.MetaError):
        await sync.sync(db_session)
    state = await db_session.get(MetaSyncState, 1)
    assert state.error.startswith("190") and state.synced_at is None


async def test_the_report_counts_the_cost_of_a_lead_and_a_client(
    client: httpx.AsyncClient, db_session: AsyncSession, meta, price_list: None
) -> None:
    await sync.sync(db_session)
    owner = await _user(db_session, UserRole.owner)
    manager = await _user(db_session, UserRole.sales_manager)
    first = await _create(client, manager)
    second = await _create(client, manager)
    await _create(client, manager)  # no creative
    for deal in (first, second):
        patched = await client.patch(f"{URL}/deals/{deal['id']}", json={"meta_ad_id": "a1"}, headers=_h(manager))
        assert patched.status_code == 200 and patched.json()["meta_ad"]["name"] == "Статика 1.4"
    deal = await db_session.get(CrmDeal, first["id"])
    deal.is_target, deal.paid_at, deal.paid_amount = True, tz.now_utc(), 70000
    await db_session.flush()

    # Correcting the creative is for the head of sales and the owner.
    fix = await client.patch(f"{URL}/deals/{first['id']}", json={"meta_ad_id": "a2"}, headers=_h(manager))
    assert fix.status_code == 403 and fix.json()["error_code"] == "ad_change_forbidden"
    unknown = await client.patch(f"{URL}/deals/{(await _create(client, manager))['id']}", json={"meta_ad_id": "x"}, headers=_h(manager))
    assert unknown.json()["error_code"] == "ad_not_found"

    period = {"from": (TODAY - timedelta(days=1)).isoformat(), "to": TODAY.isoformat()}
    report = (await client.get(f"{URL}/ads/report", params=period, headers=_h(owner))).json()
    a1 = next(r for r in report["rows"] if r["ad_id"] == "a1")
    # No rate yet: dollars only.
    assert a1["spend_usd"] == 15.5 and a1["spend"] is None and report["days_without_rate"] == 2
    assert (a1["leads"], a1["target"], a1["paid"], a1["revenue"], a1["conversations"]) == (2, 1, 1, 70000, 6)
    undefined = next(r for r in report["rows"] if r["ad_id"] is None)
    assert undefined["leads"] >= 2 and undefined["spend_usd"] == 0

    # The rate from yesterday on: ₸ = 15.50 × 500.
    rate = await client.put(f"{URL}/ads/rates", json={"valid_from": period["from"], "rate": 500}, headers=_h(owner))
    assert rate.status_code == 200 and rate.json()[0]["rate"] == 500
    report = (await client.get(f"{URL}/ads/report", params=period, headers=_h(owner))).json()
    a1 = next(r for r in report["rows"] if r["ad_id"] == "a1")
    assert a1["spend"] == 7750 and report["days_without_rate"] == 0
    assert report["status"]["enabled"] is True and report["status"]["last_day"] == TODAY.isoformat()

    # In the workbook: cost of a lead 7750 / 2, CAC 7750 / 1.
    book = load_workbook(io.BytesIO((await client.get(
        f"{URL}/analytics/export", params={"month": TODAY.strftime("%Y-%m")}, headers=_h(owner)
    )).content))
    header = [c.value for c in book["Реклама"][1]]
    line = next(r for r in book["Реклама"].iter_rows(min_row=2, values_only=True) if r[2] == "Статика 1.4")
    assert line[header.index("Цена лида, ₸")] == 3875 and line[header.index("CAC, ₸")] == 7750

    # The manager does not see the report or set the rate.
    assert (await client.get(f"{URL}/ads/report", params=period, headers=_h(manager))).status_code == 403
    assert (await client.put(f"{URL}/ads/rates", json={"valid_from": period["from"], "rate": 1}, headers=_h(manager))).status_code == 403


async def test_refresh_needs_the_token(client: httpx.AsyncClient, db_session: AsyncSession, monkeypatch) -> None:
    monkeypatch.setattr(settings, "META_ACCESS_TOKEN", "")
    owner = await _user(db_session, UserRole.owner)
    answer = await client.post(f"{URL}/ads/sync", headers=_h(owner))
    assert answer.status_code == 404 and answer.json()["error_code"] == "meta_disabled"
