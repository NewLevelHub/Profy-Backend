"""The owner's first round of CRM feedback (29.10): whose a lead becomes when
the owner accepts it, disabling staff, the list of lead sources."""

import io
from datetime import time, timedelta

import httpx
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.crm import CrmDeal, DealStage
from app.models.scheduling import Meeting, MeetingStatus, MeetingType, SlotState, SpecialistSlot
from app.models.user import User, UserRole
from app.services import tz
from tests.integration.test_crm_deals import FUTURE_DAY, URL, _create, _h, _open_slots, _user


async def test_the_owner_hands_an_accepted_lead_to_a_manager(
    client: httpx.AsyncClient, db_session: AsyncSession, price_list: None
) -> None:
    owner = await _user(db_session, UserRole.owner)
    manager = await _user(db_session, UserRole.sales_manager, "Айгерим")
    inbox = await _create(client, owner)
    path = f"{URL}/deals/{inbox['id']}/accept"

    # The owner is not a manager: without one the deal would be nobody's.
    nobody = await client.post(path, headers=_h(owner))
    assert nobody.status_code == 422 and nobody.json()["error_code"] == "manager_required"
    not_a_manager = await client.post(path, json={"manager_id": str(owner.id)}, headers=_h(owner))
    assert not_a_manager.json()["error_code"] == "manager_required"

    manager.is_active = False
    await db_session.flush()
    gone = await client.post(path, json={"manager_id": str(manager.id)}, headers=_h(owner))
    assert gone.json()["error_code"] == "manager_not_found"
    manager.is_active = True
    await db_session.flush()

    accepted = await client.post(path, json={"manager_id": str(manager.id)}, headers=_h(owner))
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["stage"] == "qualification" and accepted.json()["manager"]["id"] == str(manager.id)
    mine = await client.get(f"{URL}/deals", headers=_h(manager))
    assert inbox["id"] in {d["id"] for d in mine.json()}


async def test_a_manager_accepts_only_for_themself(
    client: httpx.AsyncClient, db_session: AsyncSession, price_list: None
) -> None:
    owner = await _user(db_session, UserRole.owner)
    manager = await _user(db_session, UserRole.sales_manager)
    other = await _user(db_session, UserRole.sales_manager)
    inbox = await _create(client, owner)
    taken = await client.post(f"{URL}/deals/{inbox['id']}/accept", json={"manager_id": str(other.id)}, headers=_h(manager))
    assert taken.status_code == 200 and taken.json()["manager"]["id"] == str(manager.id)


async def test_a_disabled_specialist_leaves_the_calendar(
    client: httpx.AsyncClient, db_session: AsyncSession, price_list: None
) -> None:
    owner = await _user(db_session, UserRole.owner)
    manager = await _user(db_session, UserRole.sales_manager)
    specialist = await _user(db_session, UserRole.psychologist, "TEST Spec")
    first = await _open_slots(db_session, specialist, FUTURE_DAY, time(10, 0), 2)
    db_session.add(Meeting(
        specialist_id=specialist.id, type=MeetingType.intro30, status=MeetingStatus.confirmed,
        start_at=first + timedelta(days=1), end_at=first + timedelta(days=1, minutes=30), handoff={},
    ))
    await db_session.flush()

    async def free() -> list:
        response = await client.get(
            "/api/v1/scheduling/free",
            params={"from": FUTURE_DAY.isoformat(), "to": FUTURE_DAY.isoformat(), "duration": 30},
            headers=_h(manager),
        )
        assert response.status_code == 200, response.text
        return response.json()["slots"]

    assert await free()
    row = next(m for m in (await client.get(f"{URL}/staff", headers=_h(owner))).json() if m["id"] == str(specialist.id))
    assert row["upcoming_meetings"] == 1

    off = await client.post(f"{URL}/staff/{specialist.id}/active", json={"active": False}, headers=_h(owner))
    assert off.status_code == 200 and off.json()["is_active"] is False
    assert await free() == []
    # Signed out at once: every request checks the account.
    assert (await client.get("/api/v1/scheduling/config", headers=_h(specialist))).status_code == 401

    on = await client.post(f"{URL}/staff/{specialist.id}/active", json={"active": True}, headers=_h(owner))
    assert on.json()["is_active"] is True and await free()

    # Not the owner themself, and only the owner.
    assert (await client.post(f"{URL}/staff/{owner.id}/active", json={"active": False}, headers=_h(owner))).status_code == 403
    assert (await client.post(f"{URL}/staff/{specialist.id}/active", json={"active": False}, headers=_h(manager))).status_code == 403


async def test_sources_are_a_list_the_owner_keeps(
    client: httpx.AsyncClient, db_session: AsyncSession, price_list: None
) -> None:
    owner = await _user(db_session, UserRole.owner)
    manager = await _user(db_session, UserRole.sales_manager)
    created = await client.post(f"{URL}/sources", json={"name": "  Реклама Meta  "}, headers=_h(owner))
    assert created.status_code == 201 and created.json()["name"] == "Реклама Meta"
    source_id = created.json()["id"]
    dup = await client.post(f"{URL}/sources", json={"name": "реклама meta"}, headers=_h(owner))
    assert dup.status_code == 409 and dup.json()["error_code"] == "source_exists"
    assert (await client.post(f"{URL}/sources", json={"name": "X"}, headers=_h(manager))).status_code == 403

    meta = (await client.get(f"{URL}/meta", headers=_h(manager))).json()
    assert source_id in {s["id"] for s in meta["sources"]}

    deal = await _create(client, manager, source={"source_id": source_id})
    assert deal["source"] == {
        "kind": "source", "id": source_id, "name": "Реклама Meta", "campaign_name": None, "adset_name": None,
        "attribution": "manual",
    }
    card = next(d for d in (await client.get(f"{URL}/deals", params={"source_id": source_id}, headers=_h(manager))).json())
    assert card["id"] == deal["id"] and card["source"]["name"] == "Реклама Meta"

    # Archived: no longer picked for new deals, kept by the ones that have it.
    archived = await client.patch(f"{URL}/sources/{source_id}", json={"active": False}, headers=_h(owner))
    assert archived.json()["active"] is False
    refused = await client.post(
        f"{URL}/deals", json={"parent": {"full_name": "Б"}, "source": {"source_id": source_id}}, headers=_h(manager)
    )
    assert refused.status_code == 422 and refused.json()["error_code"] == "source_not_found"
    kept = await client.patch(f"{URL}/deals/{deal['id']}", json={"source": {"source_id": source_id}}, headers=_h(manager))
    assert kept.status_code == 200 and kept.json()["source"]["name"] == "Реклама Meta"
    # One of the two, never both.
    both = await client.patch(
        f"{URL}/deals/{deal['id']}", json={"source": {"source_id": source_id, "ad_id": "1"}}, headers=_h(manager)
    )
    assert both.status_code == 422
    cleared = await client.patch(f"{URL}/deals/{deal['id']}", json={"source": None}, headers=_h(manager))
    assert cleared.json()["source"] is None


async def test_the_owner_dashboard(client: httpx.AsyncClient, db_session: AsyncSession, price_list: None) -> None:
    owner = await _user(db_session, UserRole.owner)
    manager = await _user(db_session, UserRole.sales_manager)
    specialist = await _user(db_session, UserRole.psychologist, "Айнура")
    first = await _open_slots(db_session, specialist, FUTURE_DAY, time(10, 0), 4)
    deal = CrmDeal(
        parent_id=(await _create(client, manager))["parent"]["id"], stage=DealStage.review_done,
        specialist_id=specialist.id, paid_amount=70000, paid_at=first - timedelta(days=3), psychologist_review=True,
        tariff_code="package", is_target=True,
    )
    db_session.add(deal)
    await db_session.flush()

    def call(offset_min: int, kind: MeetingType, state: MeetingStatus, reason: str | None = None, **extra) -> Meeting:
        start = first + timedelta(minutes=offset_min)
        meeting = Meeting(
            specialist_id=specialist.id, deal_id=deal.id, type=kind, status=state, start_at=start,
            end_at=start + timedelta(minutes=30 if kind == MeetingType.intro30 else 60),
            handoff={"student_name": "Камила"}, cancel_reason=reason, **extra,
        )
        db_session.add(meeting)
        return meeting

    moved = call(0, MeetingType.intro30, MeetingStatus.cancelled, "Клиент попросил")
    await db_session.flush()
    call(30, MeetingType.intro30, MeetingStatus.done, rescheduled_from_id=moved.id)
    call(120, MeetingType.intro30, MeetingStatus.no_show)
    call(180, MeetingType.main60, MeetingStatus.done)
    db_session.add(SpecialistSlot(specialist_id=specialist.id, start_at=first + timedelta(hours=5), state=SlotState.booked))
    await db_session.flush()

    day = FUTURE_DAY.isoformat()
    # The payment, three days before the calls, counts in a period that has it.
    since = (FUTURE_DAY - timedelta(days=5)).isoformat()
    response = await client.get(f"{URL}/dashboard", params={"from": since, "to": day}, headers=_h(owner))
    assert response.status_code == 200, response.text
    body = response.json()
    row = next(r for r in body["specialists"] if r["id"] == str(specialist.id))
    assert row["free_hours"] == 2.0 and row["busy_hours"] == 0.5
    assert (row["calls_done"], row["no_shows"], row["cancelled"]) == (2, 1, 1)
    assert row["revenue"] == 70000
    kinds = {(i["kind"], i["reason"]) for i in body["issues"] if i["deal_id"] == str(deal.id)}
    assert kinds == {("rescheduled", "Клиент попросил"), ("no_show", None)}
    review = body["payment_to_review"]
    assert review["count"] >= 1 and review["max_days"] >= 3.0

    assert (await client.get(f"{URL}/dashboard", params={"from": day, "to": day}, headers=_h(manager))).status_code == 403
    bad = await client.get(f"{URL}/dashboard", params={"from": day, "to": "2020-01-01"}, headers=_h(owner))
    assert bad.json()["error_code"] == "period_invalid"


async def test_the_analytics_workbook(client: httpx.AsyncClient, db_session: AsyncSession, price_list: None) -> None:
    head = await _user(db_session, UserRole.sales_head)
    marketer = await _user(db_session, UserRole.marketer)
    manager = await _user(db_session, UserRole.sales_manager, "Айгерим")
    await _create(client, manager)
    month = tz.today().strftime("%Y-%m")

    response = await client.get(f"{URL}/analytics/export", params={"month": month}, headers=_h(head))
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Воронка", "Источники", "Менеджеры", "Реклама"]
    funnel = book["Воронка"]
    assert funnel["A2"].value == "Неразобранное" and funnel["B2"].value >= 1
    managers = [row[0] for row in book["Менеджеры"].iter_rows(min_row=2, values_only=True)]
    assert "Айгерим" in managers

    # The marketer gets it without the managers; a manager does not get it.
    masked = load_workbook(io.BytesIO((await client.get(
        f"{URL}/analytics/export", params={"month": month}, headers=_h(marketer)
    )).content))
    assert "Менеджеры" not in masked.sheetnames
    assert (await client.get(f"{URL}/analytics/export", params={"month": month}, headers=_h(manager))).status_code == 403
    assert (await client.get(f"{URL}/analytics/export", params={"month": "2026-13"}, headers=_h(head))).status_code == 422
